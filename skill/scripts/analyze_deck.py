#!/usr/bin/env python3
"""analyze_deck.py deck.pptx [--geometry] — inventory a reference deck: per-slide layout, word count, shapes, charts,
font sizes; theme + fonts used; chart recipes. Use it on any new reference deck before writing its design-system.md.
Needs python-pptx (pip install python-pptx) — run with the project venv."""
import sys, os, re, json, collections, zipfile, tempfile
from pptx import Presentation
E = 914400
path = sys.argv[1]; geo = '--geometry' in sys.argv
prs = Presentation(path)
print(f"# {os.path.basename(path)}: {len(prs.slides)} slides, {prs.slide_width/E:.2f}x{prs.slide_height/E:.2f} in")
fonts = collections.Counter(); sizes = collections.Counter(); colors = collections.Counter(); wordsAll = []
for i, sl in enumerate(prs.slides, 1):
    kinds = collections.Counter(); charts = []; words = 0; texts = []
    for sh in sl.shapes:
        kinds[str(sh.shape_type).split('.')[-1].split(' ')[0] if sh.shape_type else 'PH'] += 1
        if sh.has_chart:
            c = sh.chart; charts.append(str(c.chart_type).split('.')[-1].split(' ')[0] + f"x{len(list(c.plots[0].series)) if len(c.plots) else 0}")
        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text.strip(); words += len(t.split()); texts.append(t[:50].replace('\n', ' / '))
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name: fonts[r.font.name] += 1
                    if r.font.size: sizes[int(r.font.size.pt)] += 1
                    try:
                        if r.font.color and r.font.color.type and r.font.color.rgb: colors[str(r.font.color.rgb)] += 1
                    except Exception: pass
        if geo:
            print(f"   {str(sh.shape_type).split('.')[-1].split(' ')[0] if sh.shape_type else 'PH':<12} {sh.left/E:5.2f},{sh.top/E:5.2f} {sh.width/E:5.2f}x{sh.height/E:5.2f}  {sh.name[:20]:<20} {sh.text_frame.text.strip()[:30].replace(chr(10),'/') if sh.has_text_frame else ''}")
    wordsAll.append(words)
    print(f"{i:>2} [{sl.slide_layout.name[:20]:<20}] words={words:<3} {dict(kinds)} charts={charts} | {texts[0] if texts else ''}")
ws = sorted(wordsAll); print(f"\nwords/slide: median {ws[len(ws)//2]}, p80 {ws[int(len(ws)*.8)]}, max {max(ws)}")
print("fonts:", fonts.most_common(8)); print("sizes:", sorted(sizes.items(), key=lambda kv: -kv[1])[:14]); print("text colors:", colors.most_common(8))
# theme + chart recipes straight from the XML
with zipfile.ZipFile(path) as z:
    th = [n for n in z.namelist() if n.startswith('ppt/theme/')]
    if th:
        x = z.read(th[0]).decode('utf8', 'ignore'); print("theme:", dict(re.findall(r'<a:(dk1|lt1|dk2|lt2|accent\d)>\s*<a:(?:srgbClr val|sysClr val="\w+" lastClr)="([0-9A-Fa-f]{6})"', x)), re.findall(r'<a:latin typeface="([^"]+)"', x)[:2])
    for n in sorted(n for n in z.namelist() if re.match(r'ppt/charts/chart\d+\.xml$', n))[:40]:
        x = z.read(n).decode('utf8', 'ignore')
        print(f"  {os.path.basename(n):<11} {re.findall(r'<c:(\w+Chart)>', x)} pts={re.findall(r'<c:ptCount val=\"(\d+)\"', x)[:1]} dPtFills={re.findall(r'<c:dPt>.*?srgbClr val=\"(\w{6})\"', x, flags=re.S)[:4]} expl={re.findall(r'explosion val=\"(\d+)\"', x)[:2]} lblPos={set(re.findall(r'dLblPos val=\"(\w+)\"', x))} legend={'<c:legend>' in x} lblSz={sorted(set(re.findall(r'<a:defRPr[^>]*sz=\"(\d+)\"', x)))[:4]}")
